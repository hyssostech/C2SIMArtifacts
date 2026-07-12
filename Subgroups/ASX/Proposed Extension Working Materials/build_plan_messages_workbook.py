"""Builds 'ASX Sample Plan Messages' in the sub-group's paired formats:
a diffable SpreadsheetML .xml (the git-friendly working copy) and a matching
.xlsx. Mirrors the tab and column conventions of the existing ASX Sample
Initialization/Order/Report Messages workbooks: a '... MessageBody Fields'
reference tab plus per-scenario tabs with columns
Model | C2SIM Object | Parent Type | Field | Type | Value | Notes.

Content instantiates the ASX plan-semantics module (ASX-PlanSemantics-Draft.ttl
v0.0.2) for the three validated walk scenarios (PlanSemanticsWalks/). Finding
tags [Q]/[!] follow the InstantiationReview convention; PC/PD/PM IDs reference
the walks' findings tables."""

from xml.sax.saxutils import escape

H = ["Model", "C2SIM Object", "Parent Type", "Field", "Type", "Value", "Notes"]
REF_H = ["Element", "Field", "Type", "Class that Defines Field",
         "Model that defines attribute", "Value", "Notes"]

SHEETS = []

# ---------------------------------------------------------------- reference tab
SHEETS.append(("Plan MessageBody Fields", REF_H, [
    ["Plan", "Root Element", "Message", "N/A", "C2SIM", "See hasC2SIMHeader", ""],
    ["hasC2SIMHeader", "C2SIMHeader", "Message", "C2SIM", "", "", "Same header fields as Order/Report workbooks"],
    ["hasMessageBody", "AutonomousPlanBody", "Message", "ASX (plan module)", "", "", "subclass of LOX PlanBody; draft v0.0.2"],
    ["", "", "", "", "", "", ""],
    ["AutonomousPlanBody", "Element", "lox:PlanBody", "ASX", "", "", "All additions optional atop a plain LOX plan"],
    ["isFromSender", "", "UUIDBase", "DomainMessageBody", "C2SIM", "", ""],
    ["isToReceiver", "", "UUIDBase", "DomainMessageBody", "C2SIM", "", ""],
    ["isToBeExecutedNow", "", "boolean", "PlanBody", "LOX", "", "MUST be false for plan-library variants (rule R14)"],
    ["hasGoal", "", "Goal (0..*)", "AutonomousPlanBody", "ASX", "", "inline Goal definition - the transport slot"],
    ["hasGoalReference", "", "UUIDBase (0..1)", "AutonomousPlanBody", "ASX", "", "the Goal this plan pursues"],
    ["hasContextCondition", "", "Condition (0..*)", "AutonomousPlanBody", "ASX", "", "plan applicability; plan pre-selection library (R14)"],
    ["hasPlanPhase", "", "PlanPhase (0..*)", "PlanBody", "LOX", "", "phases may be AutonomousPlanPhase"],
    ["", "", "", "", "", "", ""],
    ["AutonomousPlanPhase", "Element", "lox:PlanPhase", "ASX", "", "", "inherits exactly-1 trigger + exactly-1 completion condition"],
    ["hasPhaseFailureCondition", "", "PhaseFailureCondition (0..1)", "AutonomousPlanPhase", "ASX", "AnyTaskFailed | AllTasksFailed", "aggregation policy only; timeout/guard failures implicit (R4)"],
    ["hasOnFailurePhaseReference", "", "UUIDBase (0..1)", "AutonomousPlanPhase", "ASX", "", "BT Fallback; target trigger MUST be OnPhaseFailureTrigger (R7)"],
    ["hasExecutionPolicyCode", "", "PhaseExecutionPolicyCode (0..1)", "AutonomousPlanPhase", "ASX", "SequentialInOrder | ParallelAll | ParallelAny | PriorityFallback", "subphase control flow; order via hasPhasePriority"],
    ["hasRepetitionPolicyCode", "", "RepetitionPolicyCode (0..1)", "AutonomousPlanPhase", "ASX", "ExecuteOnce | RetryOnFailure | RepeatUntilCondition | MaintainContinuously", "[!] PM1: no cadence/interval slot yet (v0.0.3 delta)"],
    ["hasRepetitionUntilCondition", "", "Condition (0..1)", "AutonomousPlanPhase", "ASX", "", "closes the loop; LOX completion code closes each iteration"],
    ["hasRepetitionLimit", "", "nonNegativeInteger (0..1)", "AutonomousPlanPhase", "ASX", "", ""],
    ["hasGuardCondition", "", "Condition (0..*)", "AutonomousPlanPhase", "ASX", "", "conjunctive invariants (R12); violation fails phase (R4)"],
    ["hasPhaseTimeout", "", "Duration (0..1)", "AutonomousPlanPhase", "ASX", "", ""],
    ["hasRequiredAutonomyLevelCode", "", "AutonomyLevelCode (0..1)", "AutonomousPlanPhase", "ASX", "Teleop | ReCont | Automated | FullAuto", "graded per-phase autonomy (Q-D)"],
    ["hasPhasePriority", "", "nonNegativeInteger (0..1)", "AutonomousPlanPhase", "ASX", "", "explicit ordering; RDF is unordered"],
    ["hasPhaseProductReference", "", "UUIDBase (0..*)", "AutonomousPlanPhase", "ASX", "", "runtime data products (pre-allocated UUIDs - PC2 convention)"],
    ["hasGoalReference", "", "UUIDBase (0..1)", "AutonomousPlanPhase", "ASX", "", ""],
    ["", "", "", "", "", "", ""],
    ["PlanPhaseTrigger subclasses", "Element", "lox:PlanPhaseTrigger", "ASX", "", "", "fail-closed for consumers that cannot interpret them (R13)"],
    ["StateConditionTrigger", "hasCondition", "Condition (1)", "StateConditionTrigger", "ASX", "", "level-triggered on own world model (R11)"],
    ["CompositeTrigger", "hasSubTrigger / hasLogicalOperatorCode", "PlanPhaseTrigger (2..*) / LogicalOperatorCode (1)", "CompositeTrigger", "ASX", "AllSubTriggers | AnySubTrigger", "R9: discrete latch, state must hold"],
    ["HumanApprovalTrigger", "hasApprovingEntity / hasApprovalTimeout / hasOnTimeoutPhaseReference", "UUIDBase (1) / Duration (0..1) / UUIDBase (0..1)", "HumanApprovalTrigger", "ASX", "", "engagement-authority gate (W1/Q-K); Refuse = timeout disposition"],
    ["OnPhaseFailureTrigger", "hasTriggerPhase", "UUIDBase (1)", "OnPhaseFailureTrigger", "ASX", "", "failure dual of PriorPhaseCompletionTrigger"],
    ["ParentPhasePolicyTrigger", "(none)", "-", "ParentPhasePolicyTrigger", "ASX", "", "fills mandatory slot of policy-managed subphases"],
    ["", "", "", "", "", "", ""],
    ["Goal", "Element", "C2SIMContent", "ASX", "", "", "end state distinct from procedure"],
    ["hasUUID", "", "UUIDBase (1)", "Goal", "C2SIM", "", ""],
    ["hasAchievementCondition", "", "Condition (1..*)", "Goal", "ASX", "", "conjunctive (R12)"],
    ["hasGoalFailureCondition", "", "Condition (0..*)", "Goal", "ASX", "", "abandonment criteria"],
    ["hasGoalCommitmentCode", "", "GoalCommitmentCode (1)", "Goal", "ASX", "AchieveOnce | Maintain | AchieveThenMaintain", "Maintain = standing/persistent tasking (N2)"],
    ["hasGoalPriority", "", "nonNegativeInteger (0..1)", "Goal", "ASX", "", ""],
    ["hasDesiredEffectCode", "", "DesiredEffectCode (0..*)", "Goal", "C2SIM", "", "back-link to existing WHY vocabulary"],
    ["", "", "", "", "", "", ""],
    ["Condition", "Element", "C2SIMContent", "ASX", "", "", "typed parameters; no expression language"],
    ["hasConditionPredicateCode", "", "ConditionPredicateCode (1)", "Condition", "ASX", "12 predicates (see module)", ""],
    ["isNegated", "", "boolean (0..1)", "Condition", "ASX", "default false", "guard polarity / else-branches"],
    ["hasConditionSubjectReference / hasConditionObjectReference", "", "UUIDBase (0..1 each)", "Condition", "ASX", "", "[!] PM3: EntityType-filter semantics of object ref to be stated"],
    ["hasThresholdValue / hasThresholdDuration", "", "double / Duration (0..1 each)", "Condition", "ASX", "", "units defined per predicate"],
    ["hasUUID", "", "UUIDBase (0..1)", "Condition", "C2SIM", "", "so reports can cite the condition that fired"],
    ["", "", "", "", "", "", ""],
    ["Report contents (see report tabs)", "PlanExecutionStatusContent / PlanDeviationReportContent", "C2SIM ReportContent", "ASX", "", "", "additive - degrade gracefully"],
    ["Control verbs (TaskActionCode)", "ConfigureAutonomy, SuspendPlanExecution, ResumePlanExecution, OverrideAction, AbortPlanExecution, AchieveGoal", "TaskActionCode individuals", "ASX", "", "", "TASKFAILD added to TaskStatusCode (candidate core erratum)"],
    ["Normative rules", "R1-R14", "module ontology header", "ASX", "", "", "execution semantics OWL cannot carry; vocabulary + rules = the proposal"],
]))

# ---------------------------------------------------------------- CASEVAC scout
SHEETS.append(("CASEVAC Scout Plan", H, [
    ["C2SIM", "AutonomousPlanBody", "DomainMessageBody", "isFromSender", "UUIDBase", "(soldier UUID)", ""],
    ["C2SIM", "AutonomousPlanBody", "DomainMessageBody", "isToReceiver", "UUIDBase", "(scout UGV UUID)", ""],
    ["LOX", "AutonomousPlanBody", "PlanBody", "isToBeExecutedNow", "boolean", "true", ""],
    ["ASX", "AutonomousPlanBody", "PlanBody", "hasGoal", "Goal", "G1 (inline, below)", "goal transported inline - v0.0.2 fix for untransmittable Goal"],
    ["ASX", "Goal G1", "C2SIMContent", "hasUUID", "UUIDBase", "UUID-G1", ""],
    ["ASX", "Goal G1", "Goal", "hasGoalCommitmentCode", "GoalCommitmentCode", "AchieveOnce", ""],
    ["ASX", "Goal G1", "Goal", "hasAchievementCondition", "Condition", "EntityAtLocation(casualty, evac point, 25 m)", "mission WHY as data - unsayable in v0.0.1"],
    ["LOX", "AutonomousPlanBody", "PlanBody", "hasPlanPhase", "AutonomousPlanPhase", "S1 Verify route", ""],
    ["LOX", "Phase S1", "PlanPhase", "hasPlanPhaseTrigger", "OnOrderTrigger", "(execute-plan task UUID)", ""],
    ["LOX", "Phase S1", "PlanPhase", "hasPlanPhaseCompletionCondition", "Code", "AllTasksComplete", ""],
    ["C2SIM", "Phase S1", "PlanPhase", "hasTaskReference", "UUIDBase", "(recon move along Route-0)", ""],
    ["ASX", "Phase S1", "AutonomousPlanPhase", "hasGuardCondition", "Condition", "HazardOnRoute(Route-0), isNegated=true", "invariant: route clear; both replan causes (sensed obstacle, received IED report) surface as this guard"],
    ["ASX", "Phase S1", "AutonomousPlanPhase", "hasOnFailurePhaseReference", "UUIDBase", "UUID-S2", "BT Fallback (rule R7)"],
    ["ASX", "Phase S1", "AutonomousPlanPhase", "hasPhaseProductReference", "UUIDBase", "UUID-Route-1", "[!] PC2: pre-allocated UUID for runtime product - convention to state normatively"],
    ["ASX", "Phase S1", "AutonomousPlanPhase", "hasGoalReference", "UUIDBase", "UUID-G1", ""],
    ["LOX", "AutonomousPlanBody", "PlanBody", "hasPlanPhase", "AutonomousPlanPhase", "S2 Replan and continue", "fallback target"],
    ["ASX", "Phase S2", "PlanPhase", "hasPlanPhaseTrigger", "OnPhaseFailureTrigger", "hasTriggerPhase = UUID-S1", "fires ONLY on S1 failure (R7); if S2 succeeds, S1 counts succeeded for parent"],
    ["LOX", "Phase S2", "PlanPhase", "hasPlanPhaseCompletionCondition", "Code", "AllTasksComplete", ""],
    ["ASX", "Phase S2", "AutonomousPlanPhase", "hasGuardCondition", "Condition", "HazardOnRoute(new route), isNegated=true", ""],
    ["ASX", "Phase S2", "AutonomousPlanPhase", "hasOnFailurePhaseReference", "UUIDBase", "UUID-S3 (hold and report)", "chain = further alternatives; must be acyclic (R7)"],
    ["ASX", "Phase S2", "AutonomousPlanPhase", "hasPhaseProductReference", "UUIDBase", "UUID-Route-1", "same product slot - whichever phase verifies the route fills it"],
]))

# ---------------------------------------------------------------- CASEVAC transport
SHEETS.append(("CASEVAC Transport Plan", H, [
    ["C2SIM", "AutonomousPlanBody", "DomainMessageBody", "isFromSender", "UUIDBase", "(soldier UUID)", ""],
    ["C2SIM", "AutonomousPlanBody", "DomainMessageBody", "isToReceiver", "UUIDBase", "(transport UGV UUID)", ""],
    ["LOX", "AutonomousPlanBody", "PlanBody", "isToBeExecutedNow", "boolean", "true", ""],
    ["ASX", "AutonomousPlanBody", "PlanBody", "hasGoalReference", "UUIDBase", "UUID-G1", "same goal as scout plan (defined there)"],
    ["LOX", "AutonomousPlanBody", "PlanBody", "hasPlanPhase", "AutonomousPlanPhase", "T1 Load casualty", ""],
    ["LOX", "Phase T1", "PlanPhase", "hasPlanPhaseTrigger", "OnOrderTrigger", "(execute-plan task UUID)", ""],
    ["C2SIM", "Phase T1", "PlanPhase", "hasTaskReference", "UUIDBase", "(move to collection point; load)", ""],
    ["LOX", "AutonomousPlanBody", "PlanBody", "hasPlanPhase", "AutonomousPlanPhase", "T2 Transport via published route", ""],
    ["ASX", "Phase T2", "PlanPhase", "hasPlanPhaseTrigger", "CompositeTrigger", "AllSubTriggers (3 sub-triggers below)", "R9: discrete sub-triggers latch; state conditions must HOLD at firing"],
    ["ASX", "T2 trigger", "CompositeTrigger", "hasSubTrigger", "EventTrigger", "RouteAdvertisedEvent", "[!] PC3: the advertisement MESSAGE remains undefined (v0.0.1 X1/X3); trigger side ready"],
    ["ASX", "T2 trigger", "CompositeTrigger", "hasSubTrigger", "StateConditionTrigger", "EntityWithinRange(casualty, transport, 2 m)", "[!] PC1: false proxy for 'casualty loaded' - propose PayloadOnBoard predicate (v0.0.3)"],
    ["ASX", "T2 trigger", "CompositeTrigger", "hasSubTrigger", "StateConditionTrigger", "CommsAvailable(transport)", "cannot depart during comms outage on a stale check (R9)"],
    ["LOX", "Phase T2", "PlanPhase", "hasPlanPhaseCompletionCondition", "Code", "AllTasksComplete", ""],
    ["C2SIM", "Phase T2", "PlanPhase", "hasTaskReference", "UUIDBase", "(transport move via UUID-Route-1)", "task consumes S1's late-bound product"],
    ["ASX", "Phase T2", "AutonomousPlanPhase", "hasGoalReference", "UUIDBase", "UUID-G1", ""],
]))

# ---------------------------------------------------------------- DroneResponse
SHEETS.append(("DroneResponse Search Plan", H, [
    ["C2SIM", "AutonomousPlanBody", "DomainMessageBody", "isFromSender", "UUIDBase", "(operator UUID)", "one plan per sUAS (Task has exactly one performer); x4 identical shape"],
    ["C2SIM", "AutonomousPlanBody", "DomainMessageBody", "isToReceiver", "UUIDBase", "(sUAS-1 UUID)", ""],
    ["LOX", "AutonomousPlanBody", "PlanBody", "isToBeExecutedNow", "boolean", "true", ""],
    ["ASX", "AutonomousPlanBody", "PlanBody", "hasGoal", "Goal", "G1 Victim located", ""],
    ["ASX", "Goal G1", "Goal", "hasGoalCommitmentCode", "GoalCommitmentCode", "AchieveOnce", ""],
    ["ASX", "Goal G1", "Goal", "hasAchievementCondition", "Condition", "EstimateConfidenceAbove(victim-track, 0.7)", "achievement by BELIEF state - 'person detected with sufficient confidence'"],
    ["LOX", "AutonomousPlanBody", "PlanBody", "hasPlanPhase", "AutonomousPlanPhase", "D1 Area search (sector A)", ""],
    ["LOX", "Phase D1", "PlanPhase", "hasPlanPhaseTrigger", "OnOrderTrigger", "(execute-plan task UUID)", ""],
    ["LOX", "Phase D1", "PlanPhase", "hasPlanPhaseCompletionCondition", "Code", "AllTasksComplete", "closes each sweep ITERATION"],
    ["ASX", "Phase D1", "AutonomousPlanPhase", "hasRepetitionPolicyCode", "RepetitionPolicyCode", "RepeatUntilCondition", "loop decorator"],
    ["ASX", "Phase D1", "AutonomousPlanPhase", "hasRepetitionUntilCondition", "Condition", "EstimateConfidenceAbove(victim-track, 0.7)", "closes the LOOP - v0.0.2 split; under v0.0.1 the loop ended after one sweep"],
    ["ASX", "Phase D1", "AutonomousPlanPhase", "hasRequiredAutonomyLevelCode", "AutonomyLevelCode", "FullAuto", ""],
    ["ASX", "Phase D1", "AutonomousPlanPhase", "hasGuardCondition", "Condition", "EntityHealthBelow(self, 0.2), isNegated=true", "[!] PD2: battery proxy - propose RemainingEnduranceBelow (v0.0.3)"],
    ["ASX", "Phase D1", "AutonomousPlanPhase", "hasOnFailurePhaseReference", "UUIDBase", "UUID-D3", ""],
    ["LOX", "AutonomousPlanBody", "PlanBody", "hasPlanPhase", "AutonomousPlanPhase", "D2 Track victim", ""],
    ["ASX", "Phase D2", "PlanPhase", "hasPlanPhaseTrigger", "StateConditionTrigger", "EstimateConfidenceAbove(victim-track, 0.7)", "search-to-track transition; level-triggered (R11)"],
    ["LOX", "Phase D2", "PlanPhase", "hasPlanPhaseCompletionCondition", "Code", "OtherOrderReceived", "tracks until relieved"],
    ["LOX", "AutonomousPlanBody", "PlanBody", "hasPlanPhase", "AutonomousPlanPhase", "D3 Replan path", ""],
    ["ASX", "Phase D3", "PlanPhase", "hasPlanPhaseTrigger", "OnPhaseFailureTrigger", "hasTriggerPhase = UUID-D1", ""],
    ["ASX", "Phase D3", "AutonomousPlanPhase", "hasOnFailurePhaseReference", "UUIDBase", "UUID-D4", "chain end = escalate to human"],
    ["LOX", "AutonomousPlanBody", "PlanBody", "hasPlanPhase", "AutonomousPlanPhase", "D4 Request assistance / hold", ""],
    ["ASX", "Phase D4", "PlanPhase", "hasPlanPhaseTrigger", "OnPhaseFailureTrigger", "hasTriggerPhase = UUID-D3", ""],
    ["C2SIM", "Phase D4", "PlanPhase", "hasTaskReference", "UUIDBase", "(HoldInPlace)", "[!] PD3: hold + deviation report + await = escalation pattern to name; request rides on RequestBody"],
    ["LOX", "Phase D4", "PlanPhase", "hasPlanPhaseCompletionCondition", "Code", "OtherOrderReceived", "operator's response is the exit"],
]))

# ---------------------------------------------------------------- MDARS
SHEETS.append(("MDARS Standing Watch Plan", H, [
    ["C2SIM", "AutonomousPlanBody", "DomainMessageBody", "isFromSender", "UUIDBase", "(host console UUID)", "one guard : many platforms; per-platform plan"],
    ["C2SIM", "AutonomousPlanBody", "DomainMessageBody", "isToReceiver", "UUIDBase", "(MDARS-1 UUID)", ""],
    ["LOX", "AutonomousPlanBody", "PlanBody", "isToBeExecutedNow", "boolean", "true", ""],
    ["ASX", "AutonomousPlanBody", "PlanBody", "hasGoal", "Goal", "G1 Depot region secure", ""],
    ["ASX", "Goal G1", "Goal", "hasGoalCommitmentCode", "GoalCommitmentCode", "Maintain", "never-terminal commitment - unsayable in core tasking or v0.0.1"],
    ["ASX", "Goal G1", "Goal", "hasAchievementCondition", "Condition", "ThreatDetectedInArea(region A), isNegated=true", "maintain: no unresolved threat in region"],
    ["LOX", "AutonomousPlanBody", "PlanBody", "hasPlanPhase", "AutonomousPlanPhase", "M1 Standing patrol + inventory", ""],
    ["LOX", "Phase M1", "PlanPhase", "hasPlanPhaseTrigger", "OnOrderTrigger", "(execute-plan task UUID)", ""],
    ["LOX", "Phase M1", "PlanPhase", "hasPlanPhaseCompletionCondition", "Code", "OtherOrderReceived", "per module guidance for standing phases"],
    ["C2SIM", "Phase M1", "PlanPhase", "hasTaskReference", "UUIDBase", "(PATROL region A)", "LOX task verb"],
    ["C2SIM", "Phase M1", "PlanPhase", "hasTaskReference", "UUIDBase", "(RF-tag inventory - UseCapability)", "inventory rides in the same phase (done during patrol)"],
    ["ASX", "Phase M1", "AutonomousPlanPhase", "hasRepetitionPolicyCode", "RepetitionPolicyCode", "MaintainContinuously", "[!] PM1: revisit CADENCE has no slot - propose hasRepetitionInterval (MDARS + MCM both demand it)"],
    ["ASX", "Phase M1", "AutonomousPlanPhase", "hasGuardCondition", "Condition", "CommsAvailable(MDARS-1, host console)", "platform out of contact = exception, not routine"],
    ["ASX", "Phase M1", "AutonomousPlanPhase", "hasGoalReference", "UUIDBase", "UUID-G1", ""],
    ["LOX", "AutonomousPlanBody", "PlanBody", "hasPlanPhase", "AutonomousPlanPhase", "M2 Assess and challenge", "exception path"],
    ["ASX", "Phase M2", "PlanPhase", "hasPlanPhaseTrigger", "CompositeTrigger", "AnySubTrigger (3 sub-triggers below)", "exception = OR over the extraction's three exceptional events"],
    ["ASX", "M2 trigger", "CompositeTrigger", "hasSubTrigger", "StateConditionTrigger", "ThreatDetectedInArea(region A)", "intruder"],
    ["ASX", "M2 trigger", "CompositeTrigger", "hasSubTrigger", "StateConditionTrigger", "EntityHealthBelow(self, 0.5)", "[!] PM2: 'robot trapped' proxy - propose EntityImmobilized (v0.0.3)"],
    ["ASX", "M2 trigger", "CompositeTrigger", "hasSubTrigger", "StateConditionTrigger", "ThreatDetectedInArea(region A, type=fire)", "[!] PM3: EntityType filter semantics of object ref to be stated (ties to X4)"],
    ["C2SIM", "Phase M2", "PlanPhase", "hasTaskReference", "UUIDBase", "(assess contact; audio challenge)", "non-lethal challenge-and-response ROE"],
    ["LOX", "Phase M2", "PlanPhase", "hasPlanPhaseCompletionCondition", "Code", "AllTasksComplete", ""],
    ["ASX", "Phase M2", "AutonomousPlanPhase", "hasRequiredAutonomyLevelCode", "AutonomyLevelCode", "ReCont", "guard takes direct control on exception - autonomy drops by design"],
    ["ASX", "(supervision)", "-", "-", "TaskActionCode", "SuspendPlanExecution / ResumePlanExecution", "guard intervention loop; declarative preemption not needed for HOTL case (walk finding)"],
]))

# ---------------------------------------------------------------- status report
SHEETS.append(("Plan Execution Status Rpt", H, [
    ["C2SIM", "ReportBody", "DomainMessageBody", "isFromSender", "UUIDBase", "(scout UGV UUID)", ""],
    ["C2SIM", "ReportBody", "DomainMessageBody", "isToReceiver", "UUIDBase", "(soldier UUID)", ""],
    ["C2SIM", "ReportBody", "ReportBody", "hasReportingEntity", "UUIDBase", "(scout UGV UUID)", ""],
    ["ASX", "ReportBody", "ReportBody", "hasReportContent", "PlanExecutionStatusContent", "(below)", "additive - degrades gracefully for legacy consumers"],
    ["LOX", "PlanExecutionStatusContent", "ReportContent", "hasPlanPhaseReference", "UUIDBase", "UUID-S1", "which phase"],
    ["ASX", "PlanExecutionStatusContent", "ReportContent", "hasPhaseOutcomeCode", "PhaseOutcomeCode", "PhaseInProgress", ""],
    ["ASX", "PlanExecutionStatusContent", "ReportContent", "hasActiveGoalReference", "UUIDBase", "UUID-G1", "the system's active INTENTION (BDI transparency)"],
    ["ASX", "PlanExecutionStatusContent", "ReportContent", "hasAttemptNumber", "nonNegativeInteger", "1", "prevents premature contingency firing under retry (R10)"],
    ["ASX", "PlanExecutionStatusContent", "ReportContent", "hasEstimatedPhaseCompletionTime", "TimeInstant", "(DateTime)", "CASEVAC 'expected time to complete phased tasks' - closes old X5 residual"],
    ["C2SIM", "ReportContent", "ReportContent", "hasTimeOfObservation", "DateTime", "(timestamp)", ""],
]))

# ---------------------------------------------------------------- deviation report
SHEETS.append(("Plan Deviation Report", H, [
    ["C2SIM", "ReportBody", "DomainMessageBody", "isFromSender", "UUIDBase", "(scout UGV UUID)", ""],
    ["C2SIM", "ReportBody", "DomainMessageBody", "isToReceiver", "UUIDBase", "(soldier UUID)", ""],
    ["ASX", "ReportBody", "ReportBody", "hasReportContent", "PlanDeviationReportContent", "(below)", "closes [!] X2 recorded in 'ASX Sample Report Messages / CASEVAC Explainable'"],
    ["ASX", "PlanDeviationReportContent", "ReportContent", "hasDeviationTypeCode", "DeviationTypeCode", "FallbackActivated", "[Q] PD1/PM4: vocabulary scope (BehaviorAdjusted / ExceptionRaised) - v0.0.3 decision"],
    ["ASX", "PlanDeviationReportContent", "ReportContent", "hasDeviationCondition", "Condition", "(guard UUID-guard-route0-clear)", "the violated guard, citable by UUID"],
    ["LOX", "PlanDeviationReportContent", "ReportContent", "hasPlanPhaseReference", "UUIDBase", "UUID-S1", "affected phase"],
    ["ASX", "PlanDeviationReportContent", "ReportContent", "hasRationaleText", "string", "Suspected IED reported on Route-0 at bridge crossing; replanning via northern ford, transport clearance verified.", "the WHY - CASEVAC 'explainable reasons', structured"],
    ["SMX", "PlanDeviationReportContent", "ReportContent", "hasConfidenceLevel", "double", "0.85", "reuses SMX observation confidence"],
    ["C2SIM", "ReportContent", "ReportContent", "hasTimeOfObservation", "DateTime", "(timestamp)", ""],
    ["ASX", "(schema note)", "-", "-", "-", "-", "adapted from DroneResponse explanation schema {event, action, reasoning, change, confidence}; [!] PD4: 'action taken' slot proposed for v0.0.3"],
]))

# ================================================================ emitters

def emit_spreadsheetml(path):
    def cell(v):
        return f'    <Cell><Data ss:Type="String">{escape(str(v))}</Data></Cell>\n'
    ws_parts = []
    for name, hdr, rows in SHEETS:
        n_rows = len(rows) + 1
        n_cols = len(hdr)
        body = '   <Row>\n' + ''.join(cell(h) for h in hdr) + '   </Row>\n'
        for r in rows:
            body += '   <Row>\n' + ''.join(cell(v) for v in r) + '   </Row>\n'
        ws_parts.append(
            f' <Worksheet ss:Name="{escape(name)}">\n'
            f'  <Table ss:ExpandedColumnCount="{n_cols}" ss:ExpandedRowCount="{n_rows}"'
            f' x:FullColumns="1" x:FullRows="1" ss:DefaultRowHeight="14.5">\n'
            + body +
            '  </Table>\n'
            ' </Worksheet>\n')
    xml = (
        '<?xml version="1.0"?>\n'
        '<?mso-application progid="Excel.Sheet"?>\n'
        '<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet"\n'
        ' xmlns:o="urn:schemas-microsoft-com:office:office"\n'
        ' xmlns:x="urn:schemas-microsoft-com:office:excel"\n'
        ' xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet"\n'
        ' xmlns:html="http://www.w3.org/TR/REC-html40">\n'
        ' <DocumentProperties xmlns="urn:schemas-microsoft-com:office:office">\n'
        '  <Author>C2SIM ASX Sub-group (plan-semantics working session)</Author>\n'
        '  <Created>2026-07-12T00:00:00Z</Created>\n'
        '  <Version>16.00</Version>\n'
        ' </DocumentProperties>\n'
        ' <Styles>\n'
        '  <Style ss:ID="Default" ss:Name="Normal">\n'
        '   <Alignment ss:Vertical="Bottom"/>\n'
        '   <Borders/>\n'
        '   <Font ss:FontName="Aptos Narrow" x:Family="Swiss" ss:Size="11" ss:Color="#000000"/>\n'
        '   <Interior/>\n'
        '   <NumberFormat/>\n'
        '   <Protection/>\n'
        '  </Style>\n'
        ' </Styles>\n'
        + ''.join(ws_parts) +
        '</Workbook>\n')
    with open(path, 'w') as f:
        f.write(xml)


def emit_xlsx(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    for name, hdr, rows in SHEETS:
        ws = wb.create_sheet(title=name[:31])
        ws.append(hdr)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([str(v) for v in r])
        widths = [14, 24, 20, 30, 26, 44, 60]
        for i, w in enumerate(widths[:len(hdr)], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    wb.save(path)


if __name__ == '__main__':
    emit_spreadsheetml('ASX Sample Plan Messages.xml')
    emit_xlsx('ASX Sample Plan Messages.xlsx')
    total = sum(len(rows) for _, _, rows in SHEETS)
    print(f"wrote ASX Sample Plan Messages.xml/.xlsx: {len(SHEETS)} sheets, {total} data rows")
