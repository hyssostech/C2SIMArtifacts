"""Builds ASX-PlanSemantics-Orientation.pptx - an orientation and reading
guide for people encountering the ASX plan-semantics material for the first
time. This deck is a MAP: it points into the material on branch
asx-plan-semantics and scaffolds progressive reading paths; the technical
argument itself lives in ASX-PlanSemantics-Briefing.pptx and
Plan-Semantics-Analysis.md. Style follows build_plan_semantics_deck.py /
InstantiationReview/build_findings_deck.py."""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

NAVY   = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT = RGBColor(0x2E, 0x86, 0x9A)
GRAY   = RGBColor(0x33, 0x33, 0x33)
MGRAY  = RGBColor(0x66, 0x66, 0x66)
RED    = RGBColor(0xC0, 0x39, 0x2B)
ORANGE = RGBColor(0xC8, 0x7F, 0x0A)
GREEN  = RGBColor(0x2E, 0x7D, 0x32)
LIGHT  = RGBColor(0xEE, 0xF1, 0xF5)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide():
    return prs.slides.add_slide(BLANK)


def box(s, l, t, w, h):
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    return tb, tf


def set_run(r, text, size, color, bold=False, italic=False, font="Calibri"):
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font


def header(s, title, kicker=None):
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.28), Inches(2.2), Pt(3))
    ln.fill.solid(); ln.fill.fore_color.rgb = ACCENT; ln.line.fill.background()
    tb, tf = box(s, 0.55, 0.45, 12.2, 0.9)
    p = tf.paragraphs[0]; set_run(p.add_run(), title, 30, NAVY, bold=True)
    if kicker:
        tb2, tf2 = box(s, 0.6, 1.0, 12.0, 0.35)
        p2 = tf2.paragraphs[0]; set_run(p2.add_run(), kicker, 13, ACCENT, italic=True)


def footer(s):
    n = len(prs.slides._sldIdLst)
    tb, tf = box(s, 0.55, 7.02, 12.2, 0.35)
    p = tf.paragraphs[0]
    set_run(p.add_run(), "C2SIM ASX Sub-group   |   Plan-Semantics Orientation Guide (DRAFT)   |   July 2026", 9, MGRAY)
    tb2, tf2 = box(s, 12.4, 7.02, 0.7, 0.35)
    p2 = tf2.paragraphs[0]; p2.alignment = PP_ALIGN.RIGHT
    set_run(p2.add_run(), str(n), 9, MGRAY)


BULLET0 = chr(0x2022) + "  "   # bullet glyph; chr() keeps this source file ASCII-only
BULLET1 = chr(0x2013) + "  "   # en dash for sub-bullets


def bullets(s, items, top=1.6, left=0.7, width=12.0, height=5.2, size=15, gap=6):
    """items: dicts {t: text, tag: bold lead-in, tagcolor, lvl, b, c, i, sz, gap, noprefix}"""
    height = min(height, 6.95 - top)
    tb, tf = box(s, left, top, width, height)
    first = True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        lvl = it.get("lvl", 0)
        p.level = lvl
        p.space_after = Pt(it.get("gap", gap))
        prefix = "" if it.get("noprefix") else (BULLET0 if lvl == 0 else BULLET1)
        sz = it.get("sz", size - 2 * lvl)
        if it.get("tag"):
            set_run(p.add_run(), prefix + it["tag"] + "  ", sz,
                    it.get("tagcolor", NAVY), bold=True)
            prefix = ""
        set_run(p.add_run(), prefix + it["t"], sz,
                it.get("c", GRAY), bold=it.get("b", False), italic=it.get("i", False))
    return tb


def panel(s, l, t, w, h, fill=LIGHT, line=None):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line:
        sh.line.color.rgb = line; sh.line.width = Pt(1.25)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def panel_text(s, l, t, w, h, title, lines, tcolor=NAVY, fill=LIGHT, size=13, tsize=15):
    panel(s, l, t, w, h, fill=fill)
    tb, tf = box(s, l + 0.15, t + 0.08, w - 0.3, h - 0.16)
    p = tf.paragraphs[0]
    set_run(p.add_run(), title, tsize, tcolor, bold=True)
    for ln in lines:
        p = tf.add_paragraph(); p.space_before = Pt(3)
        if isinstance(ln, tuple):
            txt, col = ln
        else:
            txt, col = ln, GRAY
        set_run(p.add_run(), txt, size, col)


def table(s, rows, top=1.7, left=0.7, width=12.0, col_w=None, size=12, hdr_size=12.5, row_h=0.32):
    n_rows, n_cols = len(rows), len(rows[0])
    shp = s.shapes.add_table(n_rows, n_cols, Inches(left), Inches(top), Inches(width), Inches(row_h * n_rows))
    t = shp.table
    if col_w:
        for i, w in enumerate(col_w):
            t.columns[i].width = Inches(w)
    for ri, row in enumerate(rows):
        for ci, cell_txt in enumerate(row):
            cell = t.cell(ri, ci)
            cell.margin_top = Inches(0.02); cell.margin_bottom = Inches(0.02)
            cell.margin_left = Inches(0.06); cell.margin_right = Inches(0.06)
            tfc = cell.text_frame; tfc.word_wrap = True
            p = tfc.paragraphs[0]
            color = GRAY
            if isinstance(cell_txt, tuple):
                cell_txt, color = cell_txt
            if ri == 0:
                set_run(p.add_run(), cell_txt, hdr_size, WHITE, bold=True)
                cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
            else:
                set_run(p.add_run(), cell_txt, size, color)
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE if ri % 2 else LIGHT
    return t


# ---------------------------------------------------------------- 1 TITLE
s = slide()
panel(s, 0, 0, 13.333, 7.5, fill=NAVY)
tb, tf = box(s, 1.0, 2.1, 11.3, 1.8)
p = tf.paragraphs[0]
set_run(p.add_run(), "Start Here:", 40, WHITE, bold=True)
p = tf.add_paragraph()
set_run(p.add_run(), "A Guide to the ASX Plan-Semantics Material", 34, WHITE, bold=True)
tb, tf = box(s, 1.0, 4.1, 11.3, 2.2)
p = tf.paragraphs[0]
set_run(p.add_run(), "What exists, what order to read it in, and where every ID you will meet resolves", 18, LIGHT)
p = tf.add_paragraph(); p.space_before = Pt(14)
set_run(p.add_run(), "C2SIM ASX Sub-group  -  orientation and reading guide  -  DRAFT, July 2026", 14, ACCENT)
p = tf.add_paragraph(); p.space_before = Pt(4)
set_run(p.add_run(), "Companion to ASX-PlanSemantics-Briefing.pptx (the content brief) - this deck is the map, not the argument",
        12, MGRAY)

# ---------------------------------------------------------------- 2 WHAT THIS DECK IS
s = slide()
header(s, "What this deck is - and is not", "Five minutes here saves hours of unguided reading")
bullets(s, [
    {"t": "This is a MAP. It tells you what the plan-semantics material contains, which piece answers which question, and in what order to read - it argues nothing by itself.", "b": True, "sz": 15.5},
    {"t": "The argument lives elsewhere: ASX-PlanSemantics-Briefing.pptx (19 slides) makes the case; Plan-Semantics-Analysis.md carries the full reasoning.", "sz": 15},
    {"t": "The material is sizable, and nobody should read it front to back:", "sz": 15},
    {"t": "a ~700-line analysis, an ~1,100-line OWL module, three validation walks, a 19-row verified standards matrix, a decision log, sample messages, and PDG change proposals", "lvl": 1},
    {"t": "No prior exposure is assumed. The 90-second story on the next slide is the only prerequisite for every reading path in this deck.", "sz": 15},
])
panel_text(s, 0.7, 5.2, 12.0, 1.45, "How to use this deck", [
    "Read the story (next slide), then pick a path on one of the two 'Pick your path' slides.",
    ("Keep the 'Decoder ring' slide open next to whatever you read - R-numbers, PL-numbers, and Q-letters appear everywhere and each resolves to exactly one place.", ACCENT),
], size=13.5)
footer(s)

# ---------------------------------------------------------------- 3 THE STORY
s = slide()
header(s, "The story in 90 seconds", "All you need before picking a reading path")
bullets(s, [
    {"tag": "The problem.", "t": "C2SIM tasking tells a unit WHAT to do; a human crew fills in every decision the order leaves open. A robotic system has no crew - those decisions must be transmitted, or surrendered to opaque onboard behavior.", "sz": 14.5},
    {"tag": "The chassis.", "t": "The Land Operations Extension already has a Plan construct (phases, triggers, completion conditions) - a rudimentary reactive-execution model, and the natural place to build.", "sz": 14.5},
    {"tag": "The gaps.", "t": "The ASX scenario corpus keeps demanding the same missing semantics: failure as an outcome, fallback, guards over world state, loops and standing tasks, goals distinct from procedure, plan selection by context, and intention/rationale reporting.", "sz": 14.5},
    {"tag": "The proposal.", "t": "An additive OWL module (asx/plan, draft v0.0.3) importing LOX: five new trigger types, enriched plan phases, a small typed Condition vocabulary, Goals, two report contents, execution-control verbs - pinned down by 19 normative rules (R1-R19).", "sz": 14.5},
    {"tag": "The validation.", "t": "Three independent checks: end-to-end scenario walks (CASEVAC, DroneResponse, MDARS), a capability matrix verified against ten robotic/agent standards, and an adversarial review pass.", "sz": 14.5},
    {"tag": "The status.", "t": "Draft before the sub-group. The adopt-as-baseline decision is Q-W in the decision log; nothing is applied to the standard yet.", "sz": 14.5},
])
footer(s)

# ---------------------------------------------------------------- 4 MATERIAL AT A GLANCE
s = slide()
header(s, "The material at a glance", "Two layers in two directories, plus the evidence base they cite")
panel_text(s, 0.7, 1.6, 6.1, 3.6, "Layer 1 - the proposal", [
    ("Subgroups/ASX/Proposed Extension Working Materials/", ACCENT),
    "Plan-Semantics-Analysis.md - the argument (8 sections)",
    "ASX-PlanSemantics-Draft.ttl - the module itself, v0.0.3",
    "ASX-PlanSemantics-Briefing.pptx - the 19-slide content brief",
    "PlanSemanticsWalks/ - three validation walks + README",
    "ASX Sample Plan Messages.xlsx / .xml - concrete messages",
    "PDG-Change-Proposals.md - six upstream items for the PDG",
], size=13)
panel_text(s, 7.0, 1.6, 5.7, 3.6, "Layer 2 - independent verification", [
    ("Subgroups/ASX/InstantiationReview/", ACCENT),
    "PlanSemantics-Walk.md - what base C2SIM verifiably",
    "  provides vs. ten robotic/agent standards; the 19-row",
    "  capability matrix; findings PL1-PL11; 15 RDF-verified",
    "  claims; reconciliation with the module (sec 8)",
    "Issues-And-Comments-Log.md - findings as log rows",
    "  (sec 6k); decisions Q-W..Q-Z (sec 7)",
], size=13)
panel_text(s, 0.7, 5.4, 12.0, 1.0, "The evidence base both layers cite", [
    "Subgroups/ASX/LLMExperiments/ (sourced scenario corpus + reference papers) and Subgroups/ASX/ContributedScenarios/ (the group's CASEVAC). Both predate this track.",
], size=13)
bullets(s, [
    {"t": "Everything is on git branch asx-plan-semantics (pushed to origin). To browse without tooling: github.com/hyssostech/C2SIMArtifacts -> switch branch -> navigate the folders above.",
     "c": NAVY, "sz": 13},
], top=6.45, height=0.5)
footer(s)

# ---------------------------------------------------------------- 5 PATH BY TIME
s = slide()
header(s, "Pick your path - by time budget", "Each level is self-contained; every level builds on the previous one")
table(s, [
    ["Budget", "Read this", "You come away with"],
    ["10 minutes",
     "This deck through the story slide; then two slides of the briefing deck: 'The concern in one slide' (slide 2) and 'The proposed module at a glance' (slide 9)",
     "why the track exists and what is proposed, in outline"],
    ["45 minutes",
     "ASX-PlanSemantics-Briefing.pptx end to end (19 slides); finish with 'Outcome in one paragraph' in PlanSemanticsWalks/README.md",
     "the full argument with a worked example, plus where validation stands"],
    ["Half a day",
     "Plan-Semantics-Analysis.md secs 1-4 (today's constructs, BT/BDI correlation, feasibility), skim sec 5 (the module); then PlanSemantics-Walk.md secs 3-4 (capability matrix + the answer)",
     "the reasoning in depth, and how ten robotic standards compare"],
    ["Full depth",
     "The TTL side by side with analysis sec 5; the three walks in PlanSemanticsWalks/; log sec 6k; PDG-Change-Proposals.md; decisions in log sec 7",
     "reviewer-grade command of the module and its evidence trail"],
], top=1.65, col_w=[1.5, 7.3, 3.2], size=12, row_h=0.88)
bullets(s, [
    {"t": "Whatever the budget: bring a question. The next slide routes the common ones to the artifact that answers them.", "c": NAVY, "sz": 13.5},
], top=6.35, height=0.55)
footer(s)

# ---------------------------------------------------------------- 6 PATH BY QUESTION
s = slide()
header(s, "Pick your path - by question", "The seven questions people actually arrive with")
table(s, [
    ["If your question is ...", "Read ..."],
    ["Why does C2SIM need anything new here?",
     "Briefing slides 2-7 (the concern, refresher, gaps); Plan-Semantics-Analysis.md secs 2-3"],
    ["Doesn't some robotics standard already solve this?",
     "PlanSemantics-Walk.md secs 2-4: verified against BT/Nav2, PDDL, FlexBE, MAVLink, IEEE 1872, JAUS AS6062, STANAG 4586, 4D/RCS, HTN, FIPA/BDI"],
    ["What exactly is being proposed?",
     "Plan-Semantics-Analysis.md sec 5 (constructs + rules R1-R19); ASX-PlanSemantics-Draft.ttl is the module itself"],
    ["Does it survive contact with real scenarios?",
     "PlanSemanticsWalks/ - CASEVAC, DroneResponse, MDARS instantiate end to end; README has the one-paragraph outcome"],
    ["What do the messages actually look like?",
     "ASX Sample Plan Messages.xlsx (one worksheet per message); worked example: briefing slide 15 / analysis sec 5.9"],
    ["What is the group being asked to decide?",
     "Briefing final slide; decisions Q-W..Q-Z in log sec 7; PDG-Change-Proposals.md for the upstream asks"],
    ["What is still open or known-weak?",
     "Analysis sec 7 (limitations, kept deliberately visible); PlanSemantics-Walk.md sec 8 (PL5 residual); log sec 8"],
], top=1.65, col_w=[4.1, 7.9], size=12, row_h=0.62)
footer(s)

# ---------------------------------------------------------------- 7 CATALOG: PROPOSAL LAYER
s = slide()
header(s, "Artifact catalog - the proposal layer", "Subgroups/ASX/Proposed Extension Working Materials/")
table(s, [
    ["Artifact", "What it is", "Open it when ..."],
    ["Plan-Semantics-Analysis.md",
     "The full argument in 8 sections: what C2SIM says today (2), BT/BDI/HTN correlation (3), feasibility (4), the module spec with rules (5), recommendations (6), honest limitations (7), scenario base (8)",
     "you want the reasoning behind any module element"],
    ["ASX-PlanSemantics-Draft.ttl",
     "The machine-readable module (OWL/Turtle, v0.0.3): vocabulary plus the 19 normative rules as annotations; parses with rdflib",
     "you need normative detail, or want to load and check it"],
    ["ASX-PlanSemantics-Briefing.pptx",
     "The 19-slide content walkthrough, written for people who know C2SIM but not autonomy semantics",
     "first substantive contact - the default starting point"],
    ["PlanSemanticsWalks/ (README + 3 walks)",
     "End-to-end instantiations of CASEVAC, DroneResponse, MDARS in module vocabulary; produced nine deltas, applied in v0.0.3",
     "you want proof it works, or examples of constructs composing"],
    ["ASX Sample Plan Messages .xlsx / .xml",
     "Sample messages, one worksheet per message; the .xml twin is diff-able so changes review in git",
     "you want concrete message shapes in Excel"],
    ["PDG-Change-Proposals.md",
     "Six candidate core/LOX changes packaged for the PDG (TASKFAILD, handshake codes, control verbs, event report, LOX errata, plan identity)",
     "you interface with the PDG or track upstream asks"],
], top=1.6, col_w=[2.7, 6.3, 3.0], size=11.5, row_h=0.72)
footer(s)

# ---------------------------------------------------------------- 8 CATALOG: VERIFICATION LAYER
s = slide()
header(s, "Artifact catalog - the verification layer", "Subgroups/ASX/InstantiationReview/ - the independent check on the proposal")
table(s, [
    ["Where", "What you find there"],
    ["PlanSemantics-Walk.md secs 1-2",
     "What the base standard verifiably provides (checked in RDF, not from documentation), and capability exemplars from ten robotic/agent standards"],
    ["... sec 3-4",
     "The 19-row capability matrix (YES / PARTIAL / NO, each cell cites the deciding RDF entity) and the answer to the framing question"],
    ["... sec 5",
     "Findings PL1-PL11: each gap named, with severity - failure binding, loops, state conditions, goals, authority gating, lifecycle, waypoints, failsafe, mission container, LOX defects"],
    ["... sec 7",
     "Verification notes: the 15 independently re-checked claims, including the negatives"],
    ["... sec 8",
     "Reconciliation: every PL finding against module v0.0.3 - ADDRESSED / PARTIAL / open"],
    ["Issues-And-Comments-Log.md",
     "Sec 6k: the same findings as tracked log rows. Sec 7: decisions Q-W (adopt the module as baseline), Q-X, Q-Y, Q-Z - drafted, pending review. Sec 8: why remaining gaps are unfilled"],
], top=1.6, col_w=[3.0, 9.0], size=12, row_h=0.6)
panel_text(s, 0.7, 6.1, 12.0, 0.8, "", [
    ("Two verified negatives worth remembering: IEEE 1872.2 defines no plan constructs at all, and the base EventCode cannot express task failure. Checked against sources, not assumed.", ORANGE),
], size=12.5)
footer(s)

# ---------------------------------------------------------------- 9 DECODER RING
s = slide()
header(s, "Decoder ring - every ID family you will meet", "Each ID resolves to exactly one row or section; nothing is free-floating")
table(s, [
    ["ID family", "What it names", "Where it resolves"],
    ["R1 - R19", "the module's normative execution rules", "Analysis sec 5 + annotations in the TTL"],
    ["PL1 - PL11", "standards-walk findings: gaps in base C2SIM", "Walk sec 5; live statuses in log sec 6k"],
    ["Q-W .. Q-Z", "decisions this track puts to the sub-group", "Log sec 7 (Q-A..Q-V are the earlier instantiation-review decisions)"],
    ["X1 - X6", "CASEVAC findings from the earlier review, pre-module", "Log sec 6b; re-checked in CASEVAC-PlanWalk.md"],
    ["PC / PD / PM", "per-walk findings: CASEVAC / DroneResponse / MDARS", "Each walk doc; cross-walk table in PlanSemanticsWalks/README.md"],
    ["deltas 1 - 9", "cross-walk change proposals, applied in v0.0.3", "PlanSemanticsWalks/README.md"],
    ["matrix rows 1 - 19", "capability-matrix rows (sequencing ... mission container)", "Walk sec 3"],
    ["TASKFAILD, TASKACPT, TASKRJCT, TASKSUSP", "proposed TaskStatusCode additions", "PDG-Change-Proposals.md secs 1-2; analysis secs 5.1 / 5.10"],
], top=1.65, col_w=[2.6, 4.9, 4.5], size=12, row_h=0.5)
bullets(s, [
    {"t": "Rule of thumb: an unfamiliar ID in any document is a pointer, not jargon - look it up in the column on the right and you land on its definition, evidence, and status.", "c": NAVY, "sz": 13},
], top=6.35, height=0.55)
footer(s)

# ---------------------------------------------------------------- 10 TRUST
s = slide()
header(s, "How the work was hardened", "Why you can trust the claims - and how to re-check them yourself")
bullets(s, [
    {"tag": "Adversarial review, four ways:", "t": "OWL consistency against the reference ontologies; fact-check of every claim against the repo; a BT/BDI conceptual critique; scenario mining. What it caught was fixed in v0.0.2; what survived is documented in analysis sec 7, not hidden.", "sz": 14},
    {"tag": "Validation walks:", "t": "all Turtle instance blocks in the three walks parse against the module (rdflib); no structural defects found; nine vocabulary deltas, applied in v0.0.3.", "sz": 14},
    {"tag": "Independent standards pass:", "t": "a separate session re-derived essentially the same construct set from BT/Nav2, PDDL, FlexBE, MAVLink, IEEE, JAUS, and STANAG - convergent corroboration of the design, not new invention.", "sz": 14},
    {"tag": "RDF-verified claims:", "t": "15 capability claims independently re-checked against the ontologies (Walk sec 7), including the two negatives on the verification-layer slide.", "sz": 14},
])
panel_text(s, 0.7, 5.0, 12.0, 1.65, "Re-verify it yourself", [
    "Python 3 + rdflib parses ASX-PlanSemantics-Draft.ttl and the walks' instance blocks - no special tooling.",
    "Every matrix cell and log finding cites the exact RDF entity it rests on: grep the reference ontologies to challenge any row.",
    ("Decks and workbooks are generated from build_*.py scripts kept beside them - the sources of truth are text files you can diff, not binaries.", MGRAY),
], size=13)
footer(s)

# ---------------------------------------------------------------- 11 STATUS
s = slide()
header(s, "Where the ball is", "Status as of July 2026 - and what a newcomer can usefully do")
bullets(s, [
    {"tag": "Module:", "t": "TTL at v0.0.3, committed and pushed. Covers findings PL1-PL4 and PL6-PL9, plus the plan-identity half of PL10; all nine walk deltas applied (delta 8 as a documented convention rather than TTL vocabulary).", "sz": 14},
    {"tag": "Before the group:", "t": "Q-W - review and adopt the module draft as the plan-semantics baseline; Q-X (plan identity / lifecycle), Q-Y (failsafe / geofence), Q-Z (structured routes) - drafted in v0.0.3, pending review. Nothing is applied to the standard yet.", "sz": 14},
    {"tag": "Open residuals (tracked, not hidden):", "t": "inter-task data flow (PL5, extends decision Q-E); Mission container + preemption priority (the rest of PL10); reactive cross-phase preemption, commanded re-allocation, report roll-up (analysis sec 7).", "sz": 14},
    {"tag": "Upstream:", "t": "six candidate core/LOX changes packaged for the PDG in PDG-Change-Proposals.md.", "sz": 14},
])
panel_text(s, 0.7, 5.15, 12.0, 1.5, "A good first contribution", [
    "Read the briefing deck, pick the validation walk nearest your domain, and re-run its reasoning against a scenario you know well.",
    ("File whatever breaks as a finding - the method note in PlanSemanticsWalks/README.md shows the format; that is exactly how the nine v0.0.3 deltas were produced.", ACCENT),
], size=13)
footer(s)

# ---------------------------------------------------------------- 12 PRACTICAL NOTES
s = slide()
header(s, "Practical notes", "House conventions that are not obvious from the files themselves")
bullets(s, [
    {"tag": "Branch:", "t": "asx-plan-semantics on origin (github.com/hyssostech/C2SIMArtifacts). The .md files render directly on GitHub - switch branch and browse; no tooling needed.", "sz": 14.5},
    {"tag": "Generated artifacts:", "t": "each .pptx / .xlsx has a build_*.py beside it. To change a deck or workbook, edit the script and re-run it - never edit the binary; it will be overwritten.", "sz": 14.5},
    {"tag": "Diff-able twins:", "t": "workbooks ship with an .xml twin carrying the same content, so message edits are reviewable in git like any text change.", "sz": 14.5},
    {"tag": "TTL + analysis travel together:", "t": "the module lives in ASX-PlanSemantics-Draft.ttl; Plan-Semantics-Analysis.md sec 5 is its prose companion - read them side by side when reviewing.", "sz": 14.5},
    {"tag": "Raising questions:", "t": "anchor them to an artifact + ID ('R13 in the TTL', 'PL5 in the walk', 'delta 7 in the walks README') so answers land as log rows instead of getting lost in discussion.", "sz": 14.5},
])
panel_text(s, 0.7, 5.6, 12.0, 1.05, "", [
    ("Lost? Return to this deck's two 'Pick your path' slides and the 'material at a glance' map - every plan-semantics artifact is reachable from those three.", NAVY),
], size=13.5)
footer(s)

out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "ASX-PlanSemantics-Orientation.pptx")
prs.save(out)
print("saved:", out, "slides:", len(prs.slides._sldIdLst))
